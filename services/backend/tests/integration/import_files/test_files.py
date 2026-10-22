import io
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from backend_app.integrations.import_files.csv_parser import parse_csv_chunk
from backend_app.integrations.import_files.errors import ImportFileError
from backend_app.integrations.import_files.storage import LocalImportStorage
from backend_app.integrations.import_files.xlsx_parser import parse_xlsx_chunk


def test_csv_resumes_at_complete_multibyte_multiline_records(tmp_path: Path) -> None:
    path = tmp_path / "input.csv"
    path.write_bytes('\ufeff编号,描述\r\n一,"跨行\n第二行"\r\n二,说明\r\n'.encode())
    first = parse_csv_chunk(path, limit=1)
    assert first.headers == ("编号", "描述")
    assert first.rows[0].source_row_no == 2
    assert first.rows[0].values == ("一", "跨行\n第二行")
    assert first.next_row_no == 4
    second = parse_csv_chunk(
        path, next_byte_offset=first.next_byte_offset, next_row_no=first.next_row_no, limit=1
    )
    assert second.rows[0].source_row_no == 4
    assert second.rows[0].values == ("二", "说明")
    assert second.done
    replay = parse_csv_chunk(path, limit=1)
    assert replay == first


def test_csv_rejects_cursor_inside_quoted_record(tmp_path: Path) -> None:
    path = tmp_path / "input.csv"
    payload = b'name,note\na,"first\nsecond"\nb,last\n'
    path.write_bytes(payload)
    with pytest.raises(ImportFileError):
        parse_csv_chunk(path, next_byte_offset=payload.index(b"second"), next_row_no=3)


@pytest.mark.parametrize(
    "payload",
    [
        b"a,a\n1,2\n",
        b"a,b\n1\n",
        b'a\n"unfinished\n',
        b"a\n\xff\n",
        b"a\n\x00\n",
        b"a\n" + b"x" * 32768 + b"\n",
        b"a\n" + b"x\n" * 50001,
    ],
)
def test_csv_rejects_invalid_or_unbounded_input(tmp_path: Path, payload: bytes) -> None:
    path = tmp_path / "input.csv"
    path.write_bytes(payload)
    offset = 0
    row = 1
    with pytest.raises(ImportFileError):
        while True:
            parsed = parse_csv_chunk(path, next_byte_offset=offset, next_row_no=row, limit=1000)
            if parsed.done:
                break
            offset, row = parsed.next_byte_offset, parsed.next_row_no


def test_xlsx_resumes_by_sheet_and_actual_row(tmp_path: Path) -> None:
    path = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "导入数据"
    sheet.append(["编号", "描述"])
    sheet.append(["一", "测试"])
    sheet.append([None, None])
    sheet.append(["二", "结果"])
    workbook.save(path)
    workbook.close()
    first = parse_xlsx_chunk(path, sheet_name="导入数据", limit=1)
    assert first.rows[0].values == ("一", "测试")
    assert first.next_row_no == 4
    second = parse_xlsx_chunk(path, sheet_name="导入数据", next_row_no=first.next_row_no, limit=1)
    assert second.rows[0].source_row_no == 4
    assert second.rows[0].values == ("二", "结果")
    assert second.done


@pytest.mark.parametrize("kind", ["formula", "external", "macro", "columns", "inflate", "corrupt"])
def test_xlsx_rejects_unsafe_or_unbounded_input(tmp_path: Path, kind: str) -> None:
    path = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["编号"])
    sheet.append(["=1+1" if kind == "formula" else "一"])
    if kind == "columns":
        sheet.cell(2, 201, "越界")
    workbook.save(path)
    workbook.close()
    if kind in {"external", "macro", "inflate"}:
        with ZipFile(path, "a", compression=ZIP_DEFLATED) as archive:
            if kind == "external":
                archive.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
            if kind == "macro":
                archive.writestr("xl/vbaProject.bin", b"fake")
            if kind == "inflate":
                archive.writestr("xl/worksheets/inflation.xml", b" " * (101 * 1024 * 1024))
    if kind == "corrupt":
        path.write_bytes(b"PK\x03\x04broken")
    with pytest.raises(ImportFileError):
        parse_xlsx_chunk(path)


def test_storage_is_private_immutable_and_delete_is_retryable(tmp_path: Path) -> None:
    storage = LocalImportStorage(tmp_path / "objects")
    first = storage.write("batches/one/source.csv", io.BytesIO(b"a\n1\n"))
    assert storage.write(first.key, [b"a\n1\n"]) == first
    assert storage.path(first.key).stat().st_mode & 0o777 == 0o600
    assert storage.root.stat().st_mode & 0o777 == 0o700
    assert (
        storage.verify(first.key, expected_sha256=first.sha256, expected_size=first.byte_size)
        == first
    )
    with pytest.raises(ImportFileError):
        storage.write(first.key, [b"a\n2\n"])
    with storage.open(first.key) as source:
        assert source.read() == b"a\n1\n"
    with pytest.raises(ImportFileError):
        storage.write("batches/too-large", [b"large"], max_bytes=2)
    assert len(storage.list_objects()) == 1
    storage.delete(first.key)
    storage.delete(first.key)
    assert not storage.list_objects()


@pytest.mark.parametrize(
    "key", ["../outside", "/absolute", "nested/../escape", "name\\file", "a//b"]
)
def test_storage_rejects_untrusted_keys(tmp_path: Path, key: str) -> None:
    storage = LocalImportStorage(tmp_path / "objects")
    with pytest.raises(ImportFileError):
        storage.write(key, [b"data"])


def test_storage_rejects_symbolic_link_escape(tmp_path: Path) -> None:
    storage = LocalImportStorage(tmp_path / "objects")
    (storage.root / "escape").symlink_to(tmp_path)
    with pytest.raises(ImportFileError):
        storage.write("escape/file", [b"data"])


def test_storage_scan_pages_past_registered_objects_and_retries_orphan_cleanup(
    tmp_path: Path,
) -> None:
    storage = LocalImportStorage(tmp_path / "objects")
    for key in ("a/one", "a/two", "b/one", "b/two"):
        storage.write(key, [b"data"])
    first = storage.list_objects(limit=2)
    second = storage.list_objects(limit=2, after_key=first[-1].key)
    assert [entry.key for entry in first + second] == ["a/one", "a/two", "b/one", "b/two"]
    assert storage.list_objects(limit=2, after_key=second[-1].key) == ()
    # An interrupted publication leaves a private temporary object that the same scan can remove.
    orphan = storage.root / ".pending-orphan"
    orphan.write_bytes(b"incomplete")
    storage.delete(".pending-orphan")
    assert not orphan.exists()


def test_storage_rejects_low_disk_space_without_publishing_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from collections import namedtuple

    storage = LocalImportStorage(tmp_path / "objects")
    Usage = namedtuple("Usage", "total used free")
    monkeypatch.setattr(
        "backend_app.integrations.import_files.storage.shutil.disk_usage",
        lambda root: Usage(100, 100, 0),
    )
    with pytest.raises(ImportFileError):
        storage.write("source.csv", [b"input"])
    assert storage.list_objects() == ()


def test_storage_converts_publication_failure_and_can_retry(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import os

    storage = LocalImportStorage(tmp_path / "objects")
    original = os.link

    def unavailable(*args: object, **kwargs: object) -> None:
        raise OSError("sensitive private path")

    monkeypatch.setattr("backend_app.integrations.import_files.storage.os.link", unavailable)
    with pytest.raises(ImportFileError, match="^import.file_invalid$"):
        storage.write("source.csv", [b"input"])
    assert storage.list_objects() == ()
    monkeypatch.setattr("backend_app.integrations.import_files.storage.os.link", original)
    assert storage.write("source.csv", [b"input"]).byte_size == 5


def test_private_object_corruption_is_detected_by_size_and_digest(tmp_path: Path) -> None:
    storage = LocalImportStorage(tmp_path / "objects")
    stored = storage.write("source.csv", [b"original"])
    storage.path(stored.key).write_bytes(b"tampered")
    with pytest.raises(ImportFileError):
        storage.verify(stored.key, expected_sha256=stored.sha256, expected_size=stored.byte_size)


def test_concurrent_publishers_cannot_replace_a_winning_private_object(tmp_path: Path) -> None:
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier

    storage = LocalImportStorage(tmp_path / "objects")
    barrier = Barrier(2)

    def publish(data: bytes) -> bool:
        barrier.wait(timeout=3)
        try:
            storage.write("shared.csv", [data])
            return True
        except ImportFileError:
            return False

    with ThreadPoolExecutor(max_workers=2) as pool:
        result = list(pool.map(publish, (b"first", b"second")))
    assert sorted(result) == [False, True]
    with storage.open("shared.csv") as source:
        assert source.read() in {b"first", b"second"}
    assert [entry.key for entry in storage.list_objects()] == ["shared.csv"]


def test_streaming_object_scan_yields_without_materializing_all_directory_entries(
    tmp_path: Path,
) -> None:
    storage = LocalImportStorage(tmp_path / "objects")
    for number in range(250):
        storage.write(f"batch/{number:04d}.csv", [b"test"])
    scan = storage.iter_objects()
    first = [next(scan).key for _ in range(100)]
    second = [next(scan).key for _ in range(100)]
    remaining = [entry.key for entry in scan]
    assert len(first) == len(second) == 100 and len(remaining) == 50
    assert len(set(first + second + remaining)) == 250
