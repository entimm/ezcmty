import io
import tracemalloc
from pathlib import Path

import pytest
from openpyxl import Workbook

from backend_app.integrations.import_files.csv_parser import parse_csv_chunk
from backend_app.integrations.import_files.errors import ImportFileError
from backend_app.integrations.import_files.storage import MAX_SOURCE_BYTES, LocalImportStorage
from backend_app.integrations.import_files.xlsx_parser import parse_xlsx_chunk


def test_fifty_thousand_csv_rows_resume_in_five_hundred_row_chunks_with_bounded_memory(
    tmp_path: Path,
) -> None:
    path = tmp_path / "capacity.csv"
    with path.open("w", encoding="utf-8", newline="") as output:
        output.write("编号,备注\n")
        for number in range(50000):
            output.write(f'{number},"多字节中文\n跨行"\n')
    assert path.stat().st_size <= MAX_SOURCE_BYTES
    offset = 0
    row_no = 1
    count = 0
    chunks = 0
    tracemalloc.start()
    try:
        while True:
            result = parse_csv_chunk(
                path,
                next_byte_offset=offset,
                next_row_no=row_no,
                limit=500,
            )
            chunks += 1
            assert len(result.rows) <= 500
            assert result.rows[0].values[0] == str(count)
            count += len(result.rows)
            if result.done:
                break
            offset, row_no = result.next_byte_offset, result.next_row_no
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert count == 50000 and chunks == 100
    assert result.next_row_no == 100002
    # This measures parser allocations, including replay, and does not assume deployment RSS.
    assert peak < 16 * 1024 * 1024


def test_xlsx_fifty_thousand_row_and_two_hundred_column_limits_are_readable(tmp_path: Path) -> None:
    path = tmp_path / "rows.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("导入数据")
    sheet.append(["编号"])
    for number in range(50000):
        sheet.append([number])
    workbook.save(path)
    workbook.close()
    result = parse_xlsx_chunk(path, next_row_no=50001, limit=500)
    assert result.rows[0].source_row_no == 50001
    assert result.rows[0].values == (49999,)
    assert result.done and result.next_row_no == 50002

    columns = tmp_path / "columns.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("导入数据")
    sheet.append([f"字段{number}" for number in range(200)])
    sheet.append(list(range(200)))
    workbook.save(columns)
    workbook.close()
    result = parse_xlsx_chunk(columns)
    assert len(result.headers) == 200 and len(result.rows[0].values) == 200


def test_xlsx_real_multiple_chunks_preserve_all_rows_without_loading_full_sheet(
    tmp_path: Path,
) -> None:
    path = tmp_path / "chunked.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("导入数据")
    sheet.append(["编号", "备注"])
    for number in range(1500):
        sheet.append([number, "中文"])
    workbook.save(path)
    workbook.close()
    row_no = 1
    observed: list[int] = []
    tracemalloc.start()
    try:
        while True:
            result = parse_xlsx_chunk(path, sheet_name="导入数据", next_row_no=row_no, limit=500)
            observed.extend(row.values[0] for row in result.rows)
            if result.done:
                break
            row_no = result.next_row_no
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert observed == list(range(1500))
    assert peak < 32 * 1024 * 1024


def test_upload_accepts_exact_ten_mib_and_rejects_the_next_byte(tmp_path: Path) -> None:
    storage = LocalImportStorage(tmp_path / "objects")
    data = b"x" * MAX_SOURCE_BYTES
    stored = storage.write("limit.csv", io.BytesIO(data))
    assert stored.byte_size == MAX_SOURCE_BYTES
    with pytest.raises(ImportFileError, match="import.file_too_large"):
        storage.write("oversized.csv", [data, b"x"])
    assert not storage.path("oversized.csv").exists()
