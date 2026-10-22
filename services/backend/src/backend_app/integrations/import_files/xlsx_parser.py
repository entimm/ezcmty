import re
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile
from zlib import error as ZlibError

from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import fromstring, iterparse
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from openpyxl.utils.exceptions import InvalidFileException

from backend_app.integrations.import_files.errors import ImportFileError
from backend_app.integrations.import_files.parsing import (
    MAX_CELL_LENGTH,
    MAX_COLUMNS,
    MAX_ROWS,
    ParsedChunk,
    ParsedRow,
    validate_headers,
)
from backend_app.integrations.import_files.storage import MAX_SOURCE_BYTES

MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 1000
# openpyxl eagerly loads these structures, unlike read-only sheet rows.
MAX_EAGER_XML_BYTES = 8 * 1024 * 1024


def validate_xlsx(path: str | Path) -> None:
    if Path(path).stat().st_size > MAX_SOURCE_BYTES:
        raise ImportFileError("import.file_too_large")
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if (
                len(entries) > MAX_ZIP_ENTRIES
                or len(names) != len(entries)
                or sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES
                or not {"[Content_Types].xml", "xl/workbook.xml"} <= names
            ):
                raise ImportFileError()
            for entry in entries:
                name = entry.filename
                lowered = name.lower()
                if (
                    entry.flag_bits & 1
                    or name.startswith("/")
                    or ".." in name.split("/")
                    or "vbaproject" in lowered
                    or "externallink" in lowered
                    or "embeddings/" in lowered
                    or "activex/" in lowered
                    or not name.startswith("xl/worksheets/")
                    and entry.file_size > MAX_EAGER_XML_BYTES
                ):
                    raise ImportFileError()
                if name.endswith(".rels") or name == "[Content_Types].xml":
                    content = archive.read(entry)
                    root = fromstring(content)
                    if any(
                        element.attrib.get("TargetMode") == "External"
                        or "macroEnabled" in element.attrib.get("ContentType", "")
                        for element in root.iter()
                    ):
                        raise ImportFileError()
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    # Validate coordinates before openpyxl expands sparse rows into tuples.
                    with archive.open(entry) as sheet:
                        count = 0
                        for _, element in iterparse(sheet, events=("end",)):
                            tag = element.tag.rsplit("}", 1)[-1]
                            if tag == "f" or (
                                tag == "t" and len(element.text or "") > MAX_CELL_LENGTH
                            ):
                                raise ImportFileError()
                            if tag == "c":
                                count += 1
                                coordinate = re.fullmatch(
                                    r"([A-Z]{1,3})([1-9][0-9]*)", element.get("r", "")
                                )
                                if (
                                    coordinate is None
                                    or column_index_from_string(coordinate[1]) > MAX_COLUMNS
                                    or int(coordinate[2]) > MAX_ROWS + 1
                                    or count > (MAX_ROWS + 1) * MAX_COLUMNS
                                ):
                                    raise ImportFileError()
                            element.clear()
    except BadZipFile, KeyError, ValueError, OSError, DefusedXmlException, ParseError, ZlibError:
        raise ImportFileError() from None


def _cell_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is not None and not isinstance(value, (str, int, float, bool)):
        raise ImportFileError()
    if isinstance(value, str) and (len(value) > MAX_CELL_LENGTH or "\x00" in value):
        raise ImportFileError()
    return value


def parse_xlsx_chunk(
    path: str | Path,
    *,
    sheet_name: str | None = None,
    next_row_no: int = 1,
    limit: int = 500,
    expected_headers: tuple[str, ...] | None = None,
) -> ParsedChunk:
    if not 1 <= limit <= 1000 or next_row_no < 1:
        raise ImportFileError()
    validate_xlsx(path)
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        if sheet_name is None:
            if len(workbook.sheetnames) != 1:
                raise ImportFileError()
            sheet_name = workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise ImportFileError()
        sheet = workbook[sheet_name]
        # Ignore attacker-supplied dimensions; cap actual streamed rows/cells below.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        header_cells = next(iterator, ())
        if len(header_cells) > MAX_COLUMNS or any(c.data_type == "f" for c in header_cells):
            raise ImportFileError()
        headers = tuple(str(c.value) if c.value is not None else "" for c in header_cells)
        validate_headers(headers, expected_headers)
        rows: list[ParsedRow] = []
        total_rows = 0
        current_row = 1
        for cells in iterator:
            current_row += 1
            if len(cells) > MAX_COLUMNS or current_row > MAX_ROWS + 1:
                raise ImportFileError()
            if any(c.data_type in {"f", "e"} for c in cells):
                raise ImportFileError()
            values = tuple(_cell_value(c.value) for c in cells)
            if not any(v is not None and str(v).strip() for v in values):
                continue
            total_rows += 1
            if total_rows > MAX_ROWS or len(values) > len(headers):
                raise ImportFileError()
            if current_row < max(2, next_row_no):
                continue
            if len(rows) >= limit:
                return ParsedChunk(headers, tuple(rows), 0, current_row, sheet_name, False)
            padded = values + (None,) * (len(headers) - len(values))
            rows.append(ParsedRow(current_row, padded))
        if next_row_no > current_row + 1:
            raise ImportFileError()
        return ParsedChunk(headers, tuple(rows), 0, current_row + 1, sheet_name, True)
    except ImportFileError:
        raise
    except (
        ValueError,
        KeyError,
        TypeError,
        IndexError,
        OSError,
        BadZipFile,
        DefusedXmlException,
        ParseError,
        ZlibError,
        InvalidFileException,
    ):
        raise ImportFileError() from None
    finally:
        if workbook is not None:
            workbook.close()
